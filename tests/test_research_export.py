"""Step X (PRD M5 AC7): DOCX and PDF through pandoc, checked after the fact; MD always stays."""

import stat
import sys
from collections.abc import Sequence
from pathlib import Path

import pytest
from docx import Document

from app.adapters.pandoc import PandocResult, SubprocessPandoc
from app.events import MemoryEventSink
from app.research.export import (
    ExportResult,
    docx_args,
    docx_headings,
    export_report,
    pdf_args,
)

REPORT = (
    "# Titel\n\n## Eins\n\nText [1].\n\n## Zwei\n\nMehr.\n\n## Quellen\n\n[1] A. B. 2020. "
    "https://a.example.org/x (abgerufen 2026-10-01)\n\n## Anhang A — Recherche-Brief\n\n"
    "```text\n# Brief\n\n## Ausgabe\n```\n"
)
H2 = ["Eins", "Zwei", "Quellen", "Anhang A — Recherche-Brief"]


def make_docx(path: Path, headings: Sequence[str]) -> None:
    document = Document()
    document.add_heading("Titel", level=1)
    for heading in headings:
        document.add_heading(heading, level=2)
    document.save(str(path))


class FakePandoc:
    """Writes what pandoc would: a DOCX with the report's H2s, a PDF, or nothing at all."""

    def __init__(
        self,
        *,
        docx_headings: Sequence[str] = tuple(H2),
        pdf_bytes: bytes = b"%PDF-1.7\n%fake",
        code: int = 0,
        stderr: str = "",
        write: bool = True,
    ) -> None:
        self.docx_headings = docx_headings
        self.pdf_bytes = pdf_bytes
        self.code = code
        self.stderr = stderr
        self.write = write
        self.calls: list[tuple[list[str], Path]] = []

    def run(self, args: Sequence[str], *, cwd: Path) -> PandocResult:
        self.calls.append((list(args), cwd))
        out = Path(args[args.index("-o") + 1])
        if self.write and self.code == 0:
            if out.suffix == ".docx":
                make_docx(cwd / out, self.docx_headings)
            else:
                (cwd / out).write_bytes(self.pdf_bytes)
        return PandocResult(self.code, self.stderr)


def run_export(
    tmp_path: Path, fake: FakePandoc, reference: Path | None = None
) -> tuple[ExportResult, MemoryEventSink]:
    (tmp_path / "report.md").write_text(REPORT, encoding="utf-8")
    css = tmp_path / "report.css"
    css.write_text("body {}", encoding="utf-8")
    events = MemoryEventSink()
    return export_report(tmp_path, fake, css=css, reference_docx=reference, events=events), events


def test_both_formats_are_made_and_checked(tmp_path: Path) -> None:
    fake = FakePandoc()
    result, events = run_export(tmp_path, fake)
    assert (result.docx, result.pdf) == ("ok", "ok")
    assert result.ok is True
    assert (tmp_path / "report.docx").exists()
    assert (tmp_path / "report.pdf").read_bytes().startswith(b"%PDF")
    assert [e.data["format"] for e in events.of_type("export_ok")] == ["docx", "pdf"]
    assert docx_headings(tmp_path / "report.docx") == H2


def test_the_commands_are_pandoc_with_the_documented_options(tmp_path: Path) -> None:
    ref = tmp_path / "ref.docx"
    ref.write_bytes(b"x")
    fake = FakePandoc()
    run_export(tmp_path, fake, ref)
    (docx_call, docx_cwd), (pdf_call, _) = fake.calls
    assert docx_call == docx_args(Path("report.md"), Path("report.docx"), ref)
    assert docx_call[:2] == ["report.md", "--from=markdown"]
    assert f"--reference-doc={ref}" in docx_call
    assert pdf_call == pdf_args(Path("report.md"), Path("report.pdf"), tmp_path / "report.css")
    assert "--pdf-engine=weasyprint" in pdf_call
    assert f"--css={tmp_path / 'report.css'}" in pdf_call
    assert docx_cwd == tmp_path
    assert "--reference-doc" not in " ".join(pdf_args(Path("a.md"), Path("a.pdf"), Path("c.css")))


def test_without_a_reference_document_the_default_styles_apply(tmp_path: Path) -> None:
    assert not any("reference-doc" in a for a in docx_args(Path("a.md"), Path("a.docx"), None))


def test_a_missing_pandoc_fails_visibly_and_leaves_the_markdown(tmp_path: Path) -> None:
    result, events = run_export(tmp_path, FakePandoc(code=127))
    assert (result.docx, result.pdf) == ("pandoc_missing", "pandoc_missing")
    assert result.ok is False
    assert (tmp_path / "report.md").read_text(encoding="utf-8") == REPORT
    assert not (tmp_path / "report.docx").exists()
    failed = events.of_type("export_failed")
    assert [e.level for e in failed] == ["warning", "warning"]
    assert {e.data["reason"] for e in failed} == {"pandoc_missing"}


def test_a_pandoc_error_names_the_first_line_of_what_it_said(tmp_path: Path) -> None:
    result, _ = run_export(tmp_path, FakePandoc(code=65, stderr="weasyprint failed\nmore detail"))
    assert result.docx == "pandoc_failed: weasyprint failed"
    assert result.pdf == "pandoc_failed: weasyprint failed"


def test_a_docx_without_every_heading_of_the_report_is_refused_and_removed(tmp_path: Path) -> None:
    result, events = run_export(tmp_path, FakePandoc(docx_headings=["Eins", "Zwei"]))
    assert result.docx == "invalid_output: missing headings Quellen, Anhang A — Recherche-Brief"
    assert not (tmp_path / "report.docx").exists()
    assert result.pdf == "ok"
    assert any(e.data["format"] == "docx" for e in events.of_type("export_failed"))


def test_a_pdf_that_does_not_start_with_the_pdf_signature_is_refused_and_removed(
    tmp_path: Path,
) -> None:
    result, _ = run_export(tmp_path, FakePandoc(pdf_bytes=b"<html>"))
    assert result.pdf == "invalid_output: not a PDF"
    assert not (tmp_path / "report.pdf").exists()
    assert result.docx == "ok"


def test_an_output_pandoc_did_not_write_is_refused(tmp_path: Path) -> None:
    result, _ = run_export(tmp_path, FakePandoc(write=False))
    assert (result.docx, result.pdf) == ("invalid_output: no file", "invalid_output: no file")


def test_stale_exports_of_an_earlier_run_are_removed_before_pandoc_runs(tmp_path: Path) -> None:
    (tmp_path / "report.docx").write_bytes(b"old")
    (tmp_path / "report.pdf").write_bytes(b"%PDF-old")
    result, _ = run_export(tmp_path, FakePandoc(write=False))
    assert result.docx.startswith("invalid_output")
    assert not (tmp_path / "report.docx").exists()
    assert not (tmp_path / "report.pdf").exists()


def test_an_empty_output_file_is_refused(tmp_path: Path) -> None:
    class Empty(FakePandoc):
        def run(self, args: Sequence[str], *, cwd: Path) -> PandocResult:
            (cwd / args[args.index("-o") + 1]).write_bytes(b"")
            return PandocResult(0, "")

    result, _ = run_export(tmp_path, Empty())
    assert (result.docx, result.pdf) == ("invalid_output: no file", "invalid_output: no file")


def test_a_missing_report_cannot_be_exported(tmp_path: Path) -> None:
    css = tmp_path / "report.css"
    css.write_text("", encoding="utf-8")
    result = export_report(
        tmp_path, FakePandoc(), css=css, reference_docx=None, events=MemoryEventSink()
    )
    assert (result.docx, result.pdf) == ("no_report", "no_report")


# ---- the real adapter on a stand-in executable ---------------------------------------------


def script(tmp_path: Path, body: str) -> str:
    path = tmp_path / "fake-pandoc"
    path.write_text(f"#!{sys.executable}\n{body}", encoding="utf-8")
    path.chmod(path.stat().st_mode | stat.S_IEXEC)
    return str(path)


def test_the_adapter_runs_the_binary_in_the_directory_and_reports_its_exit_code(
    tmp_path: Path,
) -> None:
    binary = script(
        tmp_path,
        "import sys, pathlib\n"
        "pathlib.Path('seen.txt').write_text(' '.join(sys.argv[1:]))\n"
        "sys.stderr.write('warning: x')\n"
        "sys.exit(3)\n",
    )
    result = SubprocessPandoc(binary).run(["a.md", "-o", "a.docx"], cwd=tmp_path)
    assert result == PandocResult(3, "warning: x")
    assert (tmp_path / "seen.txt").read_text(encoding="utf-8") == "a.md -o a.docx"


def test_the_adapter_puts_the_environments_scripts_on_the_path_for_weasyprint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("PATH", "/usr/bin:/bin")  # `uv run` would put the scripts first anyway
    binary = script(
        tmp_path, "import os, pathlib\npathlib.Path('path.txt').write_text(os.environ['PATH'])\n"
    )
    SubprocessPandoc(binary).run([], cwd=tmp_path)
    first = (tmp_path / "path.txt").read_text(encoding="utf-8").split(":")[0]
    assert first == str(Path(sys.executable).parent)


def test_a_binary_that_does_not_exist_is_exit_code_127(tmp_path: Path) -> None:
    assert SubprocessPandoc(str(tmp_path / "nope")).run([], cwd=tmp_path).code == 127


def test_a_binary_that_hangs_is_stopped(tmp_path: Path) -> None:
    binary = script(tmp_path, "import time\ntime.sleep(30)\n")
    result = SubprocessPandoc(binary, timeout_s=0.3).run([], cwd=tmp_path)
    assert result.code == 124
    assert "timed out" in result.stderr


def test_the_default_stylesheet_parses_without_errors() -> None:
    import tinycss2

    css = (Path(__file__).parents[1] / "templates" / "report.css").read_text(encoding="utf-8")
    rules = tinycss2.parse_stylesheet(css, skip_comments=True, skip_whitespace=True)
    assert rules
    assert not [r for r in rules if r.type == "error"]
