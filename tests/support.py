"""Helpers shared by test modules."""

import base64
import io
from typing import Any

from app.config import Settings
from app.store.models import Note


def make_settings(**overrides: Any) -> Settings:
    """Settings that never read a developer's `.env` (init kwargs and `UDR_*` env still apply)."""
    return Settings(_env_file=None, **overrides)  # pyright: ignore[reportCallIssue]


def make_pdf(pages: list[str], title: str = "", author: str = "") -> bytes:
    """A minimal valid PDF with one Helvetica text line per page (ASCII text only)."""
    page_count = len(pages)
    font_id = 3 + 2 * page_count
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(page_count))
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {page_count} >>",
    ]
    for i, text in enumerate(pages):
        stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4 + 2 * i} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    info_ref = ""
    if title or author:
        fields = (f"/Title ({title}) " if title else "") + (
            f"/Author ({author}) " if author else ""
        )
        objects.append(f"<< {fields}>>")
        info_ref = f" /Info {len(objects)} 0 R"
    out = b"%PDF-1.4\n"
    offsets: list[int] = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R{info_ref} >>\n"
    return out + f"{trailer}startxref\n{xref}\n%%EOF\n".encode()


def make_note(**overrides: Any) -> Note:
    """A `Note` with sensible defaults for pipeline tests."""
    fields: dict[str, Any] = {
        "run_id": "run-a",
        "note_id": "n0001",
        "kind": "source",
        "stage": "fetched",
        "url": "https://example.org/a",
        "final_url": "https://example.org/a",
        "canonical_url": "https://example.org/a",
        "doi": None,
        "title": "A source",
        "content_type": "text/html",
        "via": "local",
        "body": "Some body text. " * 30,
        "pages": (),
        "word_count": 90,
        "summary": "",
        "meta": {},
        "source_tier": "unknown",
        "utility": None,
        "derivative_of": None,
        "analysis_of": None,
        "links": (),
        "extract_failed": False,
        "claims_kept": 0,
        "claims_dropped": 0,
        "created_at": "2026-10-02T08:00:00+00:00",
    }
    return Note(**{**fields, **overrides})


# A 1-page PDF, password-protected (user password "user", RC4 128-bit), made once with Ghostscript.
_ENCRYPTED_PDF = (
    "JVBERi0xLjcKJcfsj6IKJSVJbnZvY2F0aW9uOiBncyAtcSAtZE5PUEFVU0UgLWRCQVRDSCAtc0RFVklDRT1wZGZ3"
    "cml0ZSAtc093bmVyUGFzc3dvcmQ9PyAtc1VzZXJQYXNzd29yZD0/IC1kRW5jcnlwdGlvblI9MyAtZEtleUxlbmd0"
    "aD0xMjggLXNPdXRwdXRGaWxlPT8gPwo1IDAgb2JqCjw8L0xlbmd0aCA2IDAgUi9GaWx0ZXIgL0ZsYXRlRGVjb2Rl"
    "Pj4Kc3RyZWFtCgsYSLL2BhFzFPtU0ebAuHd/xmDd+4uJk7UN0EKd+VnDkz4bOKg0Votej5poAo6yAas4rn5nTqtQ"
    "g3y/EoCyRTkvvLDJfPKE4/cXjx5h515KdPWde37zHlXCAallbmRzdHJlYW0KZW5kb2JqCjYgMCBvYmoKOTEKZW5k"
    "b2JqCjQgMCBvYmoKPDwvVHlwZS9QYWdlL01lZGlhQm94IFswIDAgNjEyIDc5Ml0KL1JvdGF0ZSAwL1BhcmVudCAz"
    "IDAgUgovUmVzb3VyY2VzPDwvUHJvY1NldFsvUERGIC9UZXh0XQovRm9udCA4IDAgUgo+PgovQ29udGVudHMgNSAw"
    "IFIKPj4KZW5kb2JqCjMgMCBvYmoKPDwgL1R5cGUgL1BhZ2VzIC9LaWRzIFsKNCAwIFIKXSAvQ291bnQgMQo+Pgpl"
    "bmRvYmoKMSAwIG9iago8PC9UeXBlIC9DYXRhbG9nIC9QYWdlcyAzIDAgUgovTWV0YWRhdGEgOSAwIFIKPj4KZW5k"
    "b2JqCjggMCBvYmoKPDwvUjcKNyAwIFI+PgplbmRvYmoKNyAwIG9iago8PC9CYXNlRm9udC9IZWx2ZXRpY2EvVHlw"
    "ZS9Gb250Ci9FbmNvZGluZy9XaW5BbnNpRW5jb2RpbmcvU3VidHlwZS9UeXBlMT4+CmVuZG9iago5IDAgb2JqCjw8"
    "L1R5cGUvTWV0YWRhdGEKL1N1YnR5cGUvWE1ML0xlbmd0aCAxMTgzPj5zdHJlYW0KPD94cGFja2V0IGJlZ2luPSfv"
    "u78nIGlkPSdXNU0wTXBDZWhpSHpyZVN6TlRjemtjOWQnPz4KPD9hZG9iZS14YXAtZmlsdGVycyBlc2M9IkNSTEYi"
    "Pz4KPHg6eG1wbWV0YSB4bWxuczp4PSdhZG9iZTpuczptZXRhLycgeDp4bXB0az0nWE1QIHRvb2xraXQgMi45LjEt"
    "MTMsIGZyYW1ld29yayAxLjYnPgo8cmRmOlJERiB4bWxuczpyZGY9J2h0dHA6Ly93d3cudzMub3JnLzE5OTkvMDIv"
    "MjItcmRmLXN5bnRheC1ucyMnIHhtbG5zOmlYPSdodHRwOi8vbnMuYWRvYmUuY29tL2lYLzEuMC8nPgo8cmRmOkRl"
    "c2NyaXB0aW9uIHJkZjphYm91dD0iIiB4bWxuczpwZGY9J2h0dHA6Ly9ucy5hZG9iZS5jb20vcGRmLzEuMy8nIHBk"
    "ZjpQcm9kdWNlcj0nR1BMIEdob3N0c2NyaXB0IDEwLjAyLjEnLz4KPHJkZjpEZXNjcmlwdGlvbiByZGY6YWJvdXQ9"
    "IiIgeG1sbnM6eG1wPSdodHRwOi8vbnMuYWRvYmUuY29tL3hhcC8xLjAvJz48eG1wOk1vZGlmeURhdGU+MjAyNi0x"
    "MC0wMlQxMDo1NjozOSswMjowMDwveG1wOk1vZGlmeURhdGU+Cjx4bXA6Q3JlYXRlRGF0ZT4yMDI2LTEwLTAyVDEw"
    "OjU2OjM5KzAyOjAwPC94bXA6Q3JlYXRlRGF0ZT4KPHhtcDpDcmVhdG9yVG9vbD5Vbmtub3duQXBwbGljYXRpb248"
    "L3htcDpDcmVhdG9yVG9vbD48L3JkZjpEZXNjcmlwdGlvbj4KPHJkZjpEZXNjcmlwdGlvbiByZGY6YWJvdXQ9IiIg"
    "eG1sbnM6eGFwTU09J2h0dHA6Ly9ucy5hZG9iZS5jb20veGFwLzEuMC9tbS8nIHhhcE1NOkRvY3VtZW50SUQ9J3V1"
    "aWQ6YmFmMjQzYzEtZjY1Yi0xMWZjLTAwMDAtOThlOTRlMjU5MjBmJy8+CjxyZGY6RGVzY3JpcHRpb24gcmRmOmFi"
    "b3V0PSIiIHhtbG5zOmRjPSdodHRwOi8vcHVybC5vcmcvZGMvZWxlbWVudHMvMS4xLycgZGM6Zm9ybWF0PSdhcHBs"
    "aWNhdGlvbi9wZGYnPjxkYzp0aXRsZT48cmRmOkFsdD48cmRmOmxpIHhtbDpsYW5nPSd4LWRlZmF1bHQnPlVudGl0"
    "bGVkPC9yZGY6bGk+PC9yZGY6QWx0PjwvZGM6dGl0bGU+PC9yZGY6RGVzY3JpcHRpb24+CjwvcmRmOlJERj4KPC94"
    "OnhtcG1ldGE+CiAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAg"
    "ICAgICAgICAgICAgICAgIAogICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAgICAg"
    "ICAgICAgICAgICAgICAgICAgICAgICAKPD94cGFja2V0IGVuZD0ndyc/PgplbmRzdHJlYW0KZW5kb2JqCjIgMCBv"
    "YmoKPDwvUHJvZHVjZXIoXDAyMFwyNjZcMjA0U1wyMjFcMjE3XDM1NFwyMjZcMzYzXDMwMVwyMzNcMDIyTFwyNjI0"
    "XDIxNE5qXDMxMlwwMDFcMzY3W1gpCi9DcmVhdGlvbkRhdGUoXDAyM1wzMzRcMzcyQ1wzNDRcMzIxXDI2MlwzMjVc"
    "MjY3XDIwMFwzMTFQXDAyMFwzNjRzXDIyNVRqXDMyNlwwMjZcMzY1RU4pCi9Nb2REYXRlKFwwMjNcMzM0XDM3MkNc"
    "MzQ0XDMyMVwyNjJcMzI1XDI2N1wyMDBcMzExUFwwMjBcMzY0c1wyMjVUalwzMjZcMDI2XDM2NUVOKT4+ZW5kb2Jq"
    "CjEwIDAgb2JqCjw8L0ZpbHRlciAvU3RhbmRhcmQgL1YgMiAvTGVuZ3RoIDEyOCAvUiAzIC9QIC00IC9PICgLo4Nf"
    "iPkDiOdOVFhBJc4UK+DeJMaw03dG4HW4kXVmcSkKL1UgKCY7b5OaV9ymj0KVGOscQO9cKL9OXk51ikFkAE5W//oB"
    "CCk+PgplbmRvYmoKeHJlZgowIDExCjAwMDAwMDAwMDAgNjU1MzUgZiAKMDAwMDAwMDUzNCAwMDAwMCBuIAowMDAw"
    "MDAxOTc1IDAwMDAwIG4gCjAwMDAwMDA0NzUgMDAwMDAgbiAKMDAwMDAwMDMzNCAwMDAwMCBuIAowMDAwMDAwMTU1"
    "IDAwMDAwIG4gCjAwMDAwMDAzMTYgMDAwMDAgbiAKMDAwMDAwMDYyNyAwMDAwMCBuIAowMDAwMDAwNTk4IDAwMDAw"
    "IG4gCjAwMDAwMDA3MTYgMDAwMDAgbiAKMDAwMDAwMjI0NSAwMDAwMCBuIAp0cmFpbGVyCjw8IC9TaXplIDExIC9S"
    "b290IDEgMCBSIC9JbmZvIDIgMCBSCi9JRCBbPEI4MDcxQjg1RUZFN0E3Q0VDNEExQjM3M0Y3Q0Y0RDI2PjxCODA3"
    "MUI4NUVGRTdBN0NFQzRBMUIzNzNGN0NGNEQyNj5dCi9FbmNyeXB0IDEwIDAgUiA+PgpzdGFydHhyZWYKMjM4OAol"
    "JUVPRgo="
)


def encrypted_pdf() -> bytes:
    """A PDF that PDFium refuses to open without its password."""
    return base64.b64decode(_ENCRYPTED_PDF)


def make_docx(
    paragraphs: list[str], *, table: list[list[str]] | None = None, table_after: int | None = None
) -> bytes:
    """A DOCX of ``paragraphs``; an optional table goes in after paragraph ``table_after``."""
    from docx import Document

    document = Document()
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if table is not None and table_after == index:
            grid = document.add_table(rows=len(table), cols=len(table[0]))
            for r, row in enumerate(table):
                for c, cell in enumerate(row):
                    grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
