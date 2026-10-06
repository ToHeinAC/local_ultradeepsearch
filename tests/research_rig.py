"""Shared rig of the Phase-2 tests: a scripted model that answers by schema or prompt, and the
parts of a research run wired on fakes."""

import json
import re
import threading
import time
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from docx import Document
from support import make_settings

from app.adapters.outbound.gateway import PreparedQuery
from app.adapters.outbound.types import ScholarlyRecord, SearchHit
from app.events import EventSink, MemoryEventSink
from app.llm.fakes import CallbackTransport, reply
from app.llm.roles import build_registry
from app.llm.service import LLMService
from app.llm.types import ChatReply, ChatRequest, Endpoint
from app.pipeline.profiles import load_research_budget, load_response_formats, load_run_rules
from app.pipeline.urls import dedup_key
from app.research.draft import Drafter
from app.research.evidence import EvidenceKeys, PackBuilder
from app.research.export import PandocResult
from app.research.fixes_model import ModelFixes
from app.research.manifest import RunSettings
from app.research.markdown import h2_list
from app.store.models import NewClaim, NewSource, Note, SourceMeta
from app.store.vault import Vault
from app.templates import load_templates

HANG_SECONDS = 120
SETTINGS = make_settings()
RULES = load_run_rules(SETTINGS.config_dir)
LIGHT = load_research_budget("light", SETTINGS.config_dir)
FORMATS = load_response_formats(SETTINGS.config_dir)
TEMPLATES = load_templates([SETTINGS.templates_dir])
URLS = {Endpoint.OWN: "http://own", Endpoint.SHARED: "http://shared"}
REGISTRY = build_registry(SETTINGS)
NOW = datetime(2026, 10, 3, 9, 0, 0, tzinfo=UTC)

QUESTION = "Wie lange dauert der Rückbau eines Forschungsreaktors?"
BRIEF = (
    f"# {QUESTION}\n\n"
    "Method: 1 Klärungsrunde\n\n"
    "## Forschungsfragen\n\n"
    "1. Wie lange dauert der Rückbau?\n"
    "2. Welche Genehmigungen sind nötig?\n\n"
    "## Ausgabe\n\n"
    "- **Berichtssprache: Deutsch.** Der gesamte Bericht wird auf Deutsch verfasst.\n"
)
RUN_SETTINGS = RunSettings(
    report_language="de",
    response_format="structured",
    template_id="technische-stellungnahme",
    interview_language="de",
    tier="light",
    summarize_model=None,
)
AUTO_SETTINGS = RUN_SETTINGS.model_copy(update={"template_id": "auto"})

DRAFT: dict[str, Any] = {
    "entities": [{"name": "Forschungsreaktor", "type": "Anlage", "required_fields": ["Dauer"]}],
    "required_formats": [],
    "required_sections": [],
    "required_section_headings": [],
    "time_horizons": [],
    "time_periods": [],
    "scope_conditions": ["Deutschland"],
    "domains": ["regulation_de_eu"],
    "tier_recommendation": "light",
    "tier_rationale": "Eine begrenzte Frage mit klarer Antwort.",
    "modality": "collect",
    "levers": {
        "voice": "analyze",
        "voice_confidence": "low",
        "domain_notes": "Behördliche Quellen zuerst; Genehmigungsrecht der letzten Jahre.",
        "inference_depth": "standard",
    },
}
CLEAN_MATRIX: dict[str, Any] = {
    "rows": [
        {"phrase": "Rückbau", "items": ["Q1"], "scope_ok": True},
        {"phrase": "Forschungsreaktor", "items": ["E1"], "scope_ok": True},
    ]
}


def q(item: str, lens: str, text: str) -> dict[str, str]:
    return {"item": item, "lens": lens, "query": text}


PLAN: dict[str, Any] = {
    "queries": [
        q("Q1", "A", "Rückbau Forschungsreaktor Dauer"),
        q("Q1", "B", "decommissioning research reactor duration study"),
        q("Q1", "C", "Rückbau Forschungsreaktor Verzögerungen Kritik"),
        q("Q2", "A", "Genehmigung Rückbau Atomgesetz"),
        q("Q2", "B", "nuclear decommissioning licensing review"),
        q("Q2", "C", "Genehmigung Rückbau Probleme Klage"),
        q("E1", "A", "Forschungsreaktor Stilllegung Stand"),
        q("E1", "C", "Forschungsreaktor Rückbau gescheitert"),
        q("E1", "C", "Forschungsreaktor Stilllegung Mehrkosten"),
        q("E1", "C", "Forschungsreaktor Rückbau Gegenargumente"),
    ]
}


class FakePreparer:
    """Stands in for the gateway's `prepare_query`: nothing leaves, every call is recorded.

    ``rewrite`` maps a query to what the sanitizer would send; ``refuse`` maps a query to the
    error the gateway would raise."""

    def __init__(
        self,
        rewrite: dict[str, tuple[str, tuple[str, ...]]] | None = None,
        refuse: dict[str, Exception] | None = None,
    ) -> None:
        self.rewrite = rewrite or {}
        self.refuse = refuse or {}
        self.calls: list[tuple[str, str]] = []

    def prepare_query(self, query: str, *, step: str) -> PreparedQuery:
        self.calls.append((query, step))
        if query in self.refuse:
            raise self.refuse[query]
        sent, removed = self.rewrite.get(query, (query, ()))
        return PreparedQuery(query, sent, removed)


class ModelCrash(BaseException):
    """A crash inside a model call that no `except Exception` may swallow."""


class SearchCrash(BaseException):
    """A crash that no `except Exception` may swallow, like a power cut."""


class FakeSearcher:
    """Stands in for the gateway's searches: answers by sent query and records every call as
    `(source, sent)`. ``fail`` maps `(source, sent)` to the error to raise; ``crash_at`` raises a
    `SearchCrash` on that (1-based) call."""

    def __init__(
        self,
        web: dict[str, list[SearchHit]] | None = None,
        scholarly: dict[tuple[str, str], list[ScholarlyRecord]] | None = None,
        fail: dict[tuple[str, str], Exception] | None = None,
        crash_at: int | None = None,
        hang_at: int | None = None,
    ) -> None:
        self.hang_at = hang_at  # the n-th call sleeps "forever" (a SIGKILL test kills it there)
        self.web = web or {}
        self.scholarly = scholarly or {}
        self.fail = fail or {}
        self.crash_at = crash_at
        self.calls: list[tuple[str, str]] = []

    def _enter(self, source: str, sent: str) -> None:
        self.calls.append((source, sent))
        if self.crash_at == len(self.calls):
            raise SearchCrash(f"crash at search {len(self.calls)}")
        if self.hang_at == len(self.calls):
            print(f"HANGING search {len(self.calls)}", flush=True)
            time.sleep(HANG_SECONDS)
        if (source, sent) in self.fail:
            raise self.fail[(source, sent)]

    def search_web(
        self,
        prepared: PreparedQuery,
        *,
        step: str,
        include_domains: Sequence[str] = (),
        max_results: int = 10,
    ) -> list[SearchHit]:
        self._enter("web", prepared.sent)
        return self.web.get(prepared.sent, [])

    def search_scholarly(
        self, prepared: PreparedQuery, *, step: str, source: str, max_results: int = 10
    ) -> list[ScholarlyRecord]:
        self._enter(source, prepared.sent)
        return self.scholarly.get((source, prepared.sent), [])


SENTENCE = (
    "Der Rückbau kerntechnischer Anlagen dauert nach den vorliegenden Angaben mehrere Jahre und "
    "erfordert eine umfangreiche Genehmigung durch die zuständige Behörde"
)


def german_text(words: int, key: str = "S1") -> str:
    """About ``words`` words of German with a citation after every sentence."""
    per_sentence = len(SENTENCE.split()) + 1
    count = max(1, round(words / per_sentence))
    return " ".join(f"{SENTENCE} [{key}]." for _ in range(count))


def claim(text: str, quote: str) -> NewClaim:
    return NewClaim(text, "supports", "scope", "empirical", "", quote, (), (), None, None, "high")


def seed_note(
    vault: Vault,
    n: int,
    *,
    title: str | None = None,
    body: str | None = None,
    summary: str = "",
    claims: Sequence[NewClaim] = (),
    tier: str = "unknown",
    meta: SourceMeta | None = None,
    derivative_of: str | None = None,
    complete: bool = True,
    failed: bool = False,
) -> Note:
    """A stored, extracted and (by default) completed source of the run."""
    text = body if body is not None else f"Quelltext {n}. " * 20
    url = f"https://seed{n}.example.org/a"
    source = NewSource(
        url=url,
        final_url=url,
        canonical_url=dedup_key(url),
        doi=None,
        title=title or f"Quelle {n}",
        content_type="text/html",
        via="local",
        body=text,
        pages=(),
        word_count=len(text.split()),
        source_tier=tier,
        derivative_of=derivative_of,
        minhash=None,
        links=(),
        meta=meta or SourceMeta(),
    )
    note, _ = vault.add_source_note(source)
    vault.save_extraction(note.note_id, summary, list(claims), dropped=0, failed=failed)
    if complete:
        vault.mark_complete(note.note_id)
    found = vault.get_note(note.note_id)
    assert found is not None
    return found


@dataclass
class ResearchModels:
    """Scripted answers by schema (or prompt for free text), and a record of every call."""

    drafts: list[dict[str, Any]] = field(default_factory=lambda: [DRAFT])
    matrices: list[dict[str, Any]] = field(default_factory=lambda: [CLEAN_MATRIX])
    plans: list[dict[str, Any]] = field(default_factory=lambda: [PLAN])
    texts: dict[str, str] = field(default_factory=lambda: {})  # section heading -> answer
    condensed: list[dict[str, Any]] | None = None
    crash_on: dict[str, int] = field(default_factory=lambda: {})  # kind -> n-th call crashes
    answers: dict[str, list[dict[str, Any]]] = field(default_factory=lambda: {})  # by schema title
    hang_on: dict[str, int] = field(default_factory=lambda: {})  # kind -> n-th call sleeps
    thinks: dict[str, list[bool]] = field(default_factory=lambda: {})  # kind -> `think` per call
    models_used: dict[str, set[str]] = field(default_factory=lambda: {})  # kind -> model tags
    polishes: list[dict[str, Any]] = field(
        default_factory=lambda: [{"hunks": [], "escalations": []}]
    )
    readabilities: list[dict[str, Any]] = field(default_factory=lambda: [{"recommendations": []}])
    errors: dict[str, Exception] = field(default_factory=lambda: {})
    calls: dict[str, int] = field(default_factory=lambda: {})
    prompts: dict[str, list[str]] = field(default_factory=lambda: {})
    num_predicts: dict[str, list[int]] = field(default_factory=lambda: {})
    _lock: threading.Lock = field(default_factory=threading.Lock)

    def count(self, kind: str) -> int:
        return self.calls.get(kind, 0)

    def __call__(self, request: ChatRequest) -> ChatReply | Exception:
        kind = str(request.schema["title"]) if request.schema else "text"
        with self._lock:
            self.calls[kind] = self.calls.get(kind, 0) + 1
            number = self.calls[kind]
            self.thinks.setdefault(kind, []).append(request.think)
            self.num_predicts.setdefault(kind, []).append(request.num_predict)
            self.models_used.setdefault(kind, set()).add(request.model)
            self.prompts.setdefault(kind, []).append(
                "\n".join(m["content"] for m in request.messages)
            )
        if self.crash_on.get(kind) == number:
            raise ModelCrash(f"crash in {kind} call {number}")
        if self.hang_on.get(kind) == number:
            print(f"HANGING {kind} {number}", flush=True)
            time.sleep(HANG_SECONDS)
        if kind in self.errors:
            return self.errors[kind]
        if kind == "text":
            return reply(self._text(request))
        return reply(json.dumps(self._answer(kind, number)))

    def _text(self, request: ChatRequest) -> str:
        """A section: the scripted text of its heading, else German text of the asked length."""
        user = request.messages[-1]["content"]
        heading = re.search(r"Section(?: to write)?: (.+)", user)
        words = re.search(r"about (\d+) words", user)
        name = heading[1] if heading else ""
        if name in self.texts:
            return self.texts[name]
        return german_text(int(words[1]) if words else 50)

    def _answer(self, kind: str, number: int) -> dict[str, Any]:
        if kind in self.answers:
            scripted = self.answers[kind]
            return scripted[min(number, len(scripted)) - 1]
        if kind == "DecompositionDraft":
            return self.drafts[min(number, len(self.drafts)) - 1]
        if kind == "CoverageMatrix":
            return self.matrices[min(number, len(self.matrices)) - 1]
        if kind == "PlanDraft":
            return self.plans[min(number, len(self.plans)) - 1]
        if kind == "PolishProposal":
            return self.polishes[min(number, len(self.polishes)) - 1]
        if kind == "ReadabilityProposal":
            return self.readabilities[min(number, len(self.readabilities)) - 1]
        if kind == "CondensedEvidence":
            if self.condensed is not None:
                return self.condensed[min(number, len(self.condensed)) - 1]
            keys = list(
                dict.fromkeys(re.findall(r"\[(S\d+)\]", self.prompts["CondensedEvidence"][-1]))
            )
            return {"lines": [{"key": k, "text": f"Verdichtet zu {k}."} for k in keys]}
        raise AssertionError(f"no scripted answer for {kind}")


def llm(
    models: ResearchModels,
    events: EventSink | None = None,
    *,
    reason_num_ctx: int | None = None,
) -> LLMService:
    return LLMService(
        REGISTRY if reason_num_ctx is None else build_registry(SETTINGS, reason_num_ctx),
        URLS,
        CallbackTransport(models),
        events or MemoryEventSink(),
        timeout_s=5,
        sleep=lambda _s: None,
    )


@dataclass
class Wired:
    """The model-driven parts of Phase 2 over a vault with seeded sources."""

    vault: Vault
    ids: list[str]
    models: ResearchModels
    events: MemoryEventSink
    keys: EvidenceKeys
    packs: PackBuilder
    drafter: Drafter
    fixes: ModelFixes


def wired(
    tmp_path: Path,
    models: ResearchModels | None = None,
    *,
    notes: int = 2,
    prompt_chars: int = 40_000,
) -> Wired:
    """A vault with ``notes`` seeded sources, evidence keys, packs, a drafter and the fixes."""
    vault = Vault(tmp_path / "udr.sqlite", "run-a")
    ids = [
        seed_note(
            vault,
            n,
            title=f"Quelle {n}",
            body="Der Rückbau dauert zehn Jahre.",
            summary=f"Zusammenfassung {n}",
            claims=[claim("Der Rückbau dauert zehn Jahre.", "Der Rückbau dauert zehn Jahre.")],
        ).note_id
        for n in range(1, notes + 1)
    ]
    events = MemoryEventSink()
    models = models or ResearchModels()
    service = llm(models, events)
    keys = EvidenceKeys(tmp_path / "keys.json")
    packs = PackBuilder(vault, keys, service, RULES, events)
    drafter = Drafter(
        service, packs, keys, RULES, events, prompt_chars=prompt_chars, condense_chars=20_000
    )
    fixes = ModelFixes(
        service, packs, keys, RULES, events, prompt_chars=prompt_chars, condense_chars=20_000
    )
    return Wired(vault, ids, models, events, keys, packs, drafter, fixes)


def make_docx(path: Path, headings: Sequence[str]) -> None:
    document = Document()
    document.add_heading("Titel", level=1)
    for heading in headings:
        document.add_heading(heading, level=2)
    document.save(str(path))


class FakePandoc:
    """Writes what pandoc would: a DOCX with the report's H2s, or nothing at all."""

    def __init__(
        self,
        *,
        docx_headings: Sequence[str] | None = None,
        code: int = 0,
        stderr: str = "",
        write: bool = True,
    ) -> None:
        self.docx_headings = docx_headings
        self.code = code
        self.stderr = stderr
        self.write = write
        self.calls: list[tuple[list[str], Path]] = []

    def run(self, args: Sequence[str], *, cwd: Path) -> PandocResult:
        self.calls.append((list(args), cwd))
        out = Path(args[args.index("-o") + 1])
        if self.write and self.code == 0 and out.suffix == ".docx":
            found = self.docx_headings
            if found is None:  # like pandoc: every H2 of the report
                found = h2_list((cwd / "report.md").read_text(encoding="utf-8"))
            make_docx(cwd / out, found)
        return PandocResult(self.code, self.stderr)
